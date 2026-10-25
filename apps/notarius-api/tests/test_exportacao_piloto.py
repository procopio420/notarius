"""Exact DOCX and normalized PDF content checks with entirely fictional text."""
from io import BytesIO
from types import SimpleNamespace
import re

from django.test import SimpleTestCase
from django.utils import timezone
from django.test import override_settings
from datetime import datetime, timezone as datetime_timezone
from docx import Document
from pypdf import PdfReader

from apps.processos.exportacao import word, pdf, descricao


def amostra():
    text = '\n'.join([
        'EXERCÍCIO FICTÍCIO SEM VALIDADE PROFISSIONAL',
        'PARTE VENDEDORA: PESSOA FICTÍCIA ÁGUA',
        'PARTE COMPRADORA: PESSOA FICTÍCIA AÇÚCAR',
        'IMÓVEL FICTÍCIO: REFERÊNCIA 000123-A',
        'PREÇO FICTÍCIO: R$ 123.456,78',
        'PAGAMENTO PREVISTO: condição fictícia, sem declaração de quitação.',
        'TEXTO LITERAL: cláusula <teste> & revisão; número 0007; ação, órgão, matrícula.',
        '',
        *[f'Linha fictícia {i:03d}: referência {i:06d}; São João; condição para revisão.'
          for i in range(1, 66)],
        'FIM DO EXERCÍCIO FICTÍCIO',
    ])
    return SimpleNamespace(corpo=text, versao=2, tipo='editada', created_at=timezone.now(),
                           modelo=SimpleNamespace(versao=1))


class ExportacaoPilotoTests(SimpleTestCase):
    @override_settings(TIME_ZONE='America/Cuiaba')
    def test_export_metadata_uses_configured_local_timezone(self):
        draft = amostra()
        draft.created_at = datetime(2026, 10, 25, 17, 0, tzinfo=datetime_timezone.utc)
        self.assertIn('25/10/2026 13:00', descricao(draft))

    def test_docx_preserves_every_saved_line_blank_and_literal_character(self):
        draft = amostra()
        document = Document(BytesIO(word(draft)))
        self.assertEqual('\n'.join(p.text for p in document.paragraphs[3:]), draft.corpo)
        self.assertAlmostEqual(document.sections[0].page_width.cm, 21, places=2)
        self.assertAlmostEqual(document.sections[0].page_height.cm, 29.7, places=1)
        self.assertFalse(document.styles['Title'].element.xpath('./w:pPr/w:pBdr'))
        self.assertFalse(document.paragraphs[0]._p.xpath('./w:pPr/w:pBdr'))

    def test_pdf_keeps_numbered_heading_on_same_page_as_following_content(self):
        draft = amostra()
        draft.corpo = '\n'.join(['Linha de contexto inteiramente fictícia.'] * 28 +
                                ['4.- TÍTULO FICTÍCIO', 'CONTEÚDO VINCULADO AO TÍTULO FICTÍCIO'] +
                                ['Continuação inteiramente fictícia.'] * 20)
        pages = [p.extract_text() for p in PdfReader(BytesIO(pdf(draft))).pages]
        title_page = next(p for p in pages if '4.- TÍTULO FICTÍCIO' in p)
        self.assertIn('CONTEÚDO VINCULADO AO TÍTULO FICTÍCIO', title_page)

    def test_paginated_pdf_keeps_entire_content_and_order_after_whitespace_normalization(self):
        draft = amostra()
        reader = PdfReader(BytesIO(pdf(draft)))
        self.assertGreater(len(reader.pages), 1)
        extracted = '\n'.join(page.extract_text() for page in reader.pages)
        for item in ('Minuta de compra e venda',
                     'Para revisão. Não constitui documento assinado ou ato concluído.',
                     descricao(draft), 'Minuta para revisão'):
            extracted = extracted.replace(item, '')
        extracted = re.sub(r'Página\s+\d+', '', extracted)
        compact = lambda text: re.sub(r'\s+', '', text)
        self.assertEqual(compact(extracted), compact(draft.corpo))
        self.assertIn('FIM DO EXERCÍCIO FICTÍCIO', reader.pages[-1].extract_text())
