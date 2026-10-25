"""Full technical rehearsal only; all approvals are simulated in an isolated DB."""
from io import BytesIO
from uuid import uuid4
from unittest.mock import patch
from django.core.files.uploadedfile import SimpleUploadedFile
from rest_framework.test import APITestCase
from apps.processos.models import Atendimento, AnexoAtendimento
from apps.processos.modelo_complementar import referencia_complementar
from apps.processos.redacao_estruturada import TOKEN
from tests import test_cenarios_homologacao as base


class ValidacaoIntegralLocalTests(APITestCase):
    setUp = base.CenariosHomologacaoTests.setUp
    post = base.CenariosHomologacaoTests.post
    seed = base.CenariosHomologacaoTests.seed
    simulate_review = base.CenariosHomologacaoTests.simulate_review
    generate = base.CenariosHomologacaoTests.generate

    def source(self, manifest, alias, values, uncertain=False):
        case = manifest['atendimento_id']
        text = 'FONTE FICTÍCIA DE TESTE — NÃO É DOCUMENTO AUTÊNTICO\n' + '\n'.join(f"{k.replace('_', ' ')}: {v}" for k, v in values.items())
        uploaded = self.client.post(f'/api/v1/atendimentos/{case}/anexos/',
            {'arquivo': SimpleUploadedFile(f'FICTICIO-INTEGRAL-{alias}.txt', text.encode('utf-8'))}, format='multipart')
        self.assertEqual(uploaded.status_code, 201, uploaded.data)
        self.post(case, 'propostas-dados', {'operacao': 'extrair', 'anexo_id': uploaded.data['id'], 'leitura_versao': 1, 'client_id': str(uuid4())})
        proposals = [p for p in self.client.get(f'/api/v1/atendimentos/{case}/propostas-dados/').data['propostas']
                     if p['origem'].get('anexo_id') == uploaded.data['id']]
        self.assertEqual(len(proposals), len(values))
        fields = []
        uncertain_proposal = None
        for key, value in values.items():
            proposal = next(p for p in proposals if p['rotulo'] == key.replace('_', ' '))
            ambiguous = uncertain and key == 'conteudo'
            if ambiguous:
                uncertain_proposal = proposal
            fields.append({'chave': key, 'valor': value, 'valor_original': proposal['valor_original'],
                'situacao': 'divergente' if ambiguous else 'confirmado', 'origem': 'documento', 'versao_esperada': 0,
                'confirmar': not ambiguous, 'conferir_fonte': not ambiguous, 'proposta_id': proposal['id'],
                'versao_proposta': 0, 'conferir_associacao': True, 'pagina_id': proposal['origem']['pagina_id'],
                'leitura_versao': 1, 'trecho_origem': proposal['trecho']})
        self.post(case, 'ficha-detalhada', {'client_id': str(uuid4()), 'modelo_id': manifest['modelo_id'],
            'operacao': 'salvar', 'objeto_id': manifest['registros'][alias], 'campos': fields,
            'justificativa': 'RESPOSTAS SIMULADAS: conferência técnica de fonte sintética no banco de teste'})
        return uploaded.data['id'], uncertain_proposal

    def review_pages(self, case):
        issues = self.post(case, 'pendencias-ficha', {'operacao': 'atualizar'})['pendencias']
        for issue in issues:
            if issue['tipo'] != 'pagina' or issue['situacao'] == 'resolvida':
                continue
            if any(p['situacao'] not in ['aceita', 'rejeitada', 'substituida'] for p in issue['contexto']['propostas']):
                continue
            self.post(case, 'pendencias-ficha', {'operacao': 'resolver', 'client_id': str(uuid4()), 'pendencia_id': issue['id'],
                'assinatura': issue['assinatura'], 'versao_esperada': issue['versao'], 'acao': 'vincular',
                'revisoes_ids': issue['contexto']['usos_atuais'], 'confirmei_revisao': True,
                'justificativa': 'REVISÃO DE PÁGINA SIMULADA: unidade TXT fictícia, sem autenticação'})

    def full_case(self, variant):
        # The older rehearsal deliberately uses invalid CPF placeholders and a
        # price with prose. A successful export now needs comparable synthetic
        # formats; changing this isolated fixture does not relax the comparator
        # or turn the local homologation scenarios into approved case data.
        from apps.processos.cenarios_homologacao import especificacao
        def formatted_fixture(selected):
            spec = especificacao(selected)
            counter = 0
            for node in spec['registros']:
                if node['tipo'] == 'pessoa':
                    counter += 1
                    digits = f'90000000{counter}'
                    for size in [10, 11]:
                        digit = (sum(int(n) * (size-i) for i, n in enumerate(digits)) * 10) % 11
                        digits += str(digit if digit < 10 else 0)
                    node['campos']['cpf'] = digits
                elif node['tipo'] == 'negociacao':
                    node['campos']['preco'] = 'R$ 200.000,00'
            return spec
        with patch('apps.processos.cenarios_homologacao.especificacao', side_effect=formatted_fixture):
            manifest = self.seed(variant)
        self.simulate_review(manifest)
        case = manifest['atendimento_id']
        reference = referencia_complementar(variant)['modelo']
        draft_model = self.post(case, 'modelo-minuta', {**reference, 'versao_esperada': 2,
            'justificativa': 'MODELO COMPLEMENTAR FICTÍCIO PARA TESTE, SEM HOMOLOGAÇÃO'})
        manifest['modelo_id'] = draft_model['id']
        uncertain_proposal = None; uncertain_file = None
        for position in reference['registros_minuta']:
            alias = position['chave']
            if alias in manifest['registros']:
                continue
            links = position['vinculos']
            scope = 'pessoa_imovel' if len(links) == 2 else 'pessoa' if 'pessoa_a' in links else 'imovel' if 'imovel' in links else 'atendimento'
            created = self.post(case, 'ficha-detalhada', {'operacao': 'criar', 'client_id': str(uuid4()), 'modelo_id': manifest['modelo_id'],
                'tipo': position['tipo'], **({'escopo': scope} if position['tipo'] == 'complemento' else {}),
                **{f'{key}_id': manifest['registros'][target] for key, target in links.items()},
                'justificativa': 'CRIAÇÃO SIMULADA DE REGISTRO FICTÍCIO PARA VALIDAÇÃO TÉCNICA'})
            manifest['registros'][alias] = created['objeto_id']
            if position['tipo'] == 'ato':
                values = {'abertura': 'ABERTURA SINTÉTICA DE TESTE', 'data_informada': 'DATA FICTÍCIA SEM VALIDADE',
                    'local': 'LOCAL INEXISTENTE DE TESTE', 'serventia': 'SERVENTIA FICTÍCIA',
                    'encerramento': 'MINUTA FICTÍCIA PARA REVISÃO, SEM ASSINATURA'}
            else:
                category = position['categoria']
                state = 'nao_realizado' if category == 'consulta' else 'previsto' if category in ['tributo', 'despesas'] else 'informado'
                content = (f'CONSULTA NÃO REALIZADA NO TESTE FICTÍCIO: {alias}' if category == 'consulta' else
                           f'PREVISÃO FICTÍCIA SEM RECOLHIMENTO: {alias}' if category in ['tributo', 'despesas'] else
                           f'CONTEÚDO SINTÉTICO SEM VALIDADE PROFISSIONAL: {alias}')
                values = {'categoria': category, 'situacao': state, 'referencia': f'REFERENCIA-FICTICIA-{alias}', 'conteudo': content}
            file_id, proposal = self.source(manifest, alias, values, uncertain=alias == 'consulta_comprador')
            if proposal:
                uncertain_proposal = proposal; uncertain_file = file_id
        self.source(manifest, 'negocio', {'quitacao': 'SEM DECLARAÇÃO DE QUITAÇÃO NO CENÁRIO FICTÍCIO'})
        self.review_pages(case)
        # A complete body with a literal source still cannot bypass an unresolved
        # interpretation or the explicit model approval, even in the rehearsal.
        self.assertEqual(self.generate(case).status_code, 400)
        approved = self.post(case, 'modelo-minuta', {**reference, 'aprovado': True, 'versao_esperada': 3,
            'justificativa': 'APROVAÇÃO EXCLUSIVAMENTE SIMULADA NO BANCO DE TESTE, NÃO PROFISSIONAL'})
        manifest['modelo_id'] = approved['id']
        state = self.client.get(f'/api/v1/atendimentos/{case}/vinculacao-minuta/').data
        self.post(case, 'vinculacao-minuta', {'client_id': str(uuid4()), 'modelo_id': approved['id'],
            'versao_esperada': 1, 'escopo_assinatura': state['escopo_assinatura'], 'registros': manifest['registros'],
            'conferir': True, 'motivo': 'VINCULAÇÃO SIMULADA DOS REGISTROS COMPLEMENTARES'})
        self.assertEqual(self.generate(case).status_code, 400)
        record = Atendimento.objects.get(pk=case)
        self.assertFalse(record.minutas_preparadas.exists())
        issue_key = f"dado:{manifest['registros']['consulta_comprador']}:conteudo"
        self.assertIn(issue_key, [p['chave'] for p in self.client.get(f'/api/v1/atendimentos/{case}/conferencia/').data['pendencias']])
        p = uncertain_proposal
        self.post(case, 'ficha-detalhada', {'client_id': str(uuid4()), 'operacao': 'salvar', 'modelo_id': approved['id'],
            'objeto_id': manifest['registros']['consulta_comprador'], 'campos': [{
                'chave': 'conteudo', 'valor': p['valor_original'].lstrip(), 'valor_original': p['valor_original'],
                'situacao': 'confirmado', 'origem': 'documento', 'versao_esperada': 1, 'confirmar': True,
                'conferir_fonte': True, 'conferir_associacao': True, 'proposta_id': p['id'], 'versao_proposta': 1,
                'pagina_id': p['origem']['pagina_id'], 'leitura_versao': 1, 'trecho_origem': p['trecho']}],
            'justificativa': 'DECISÃO HUMANA SIMULADA: dúvida de interpretação resolvida apenas para o teste'})
        self.review_pages(case)
        # A comparison opened while interpretation was divergent needs its own
        # explicit decision after field correction. Confirmation alone cannot
        # silently close the evidence review introduced in stage 5.
        issues = self.post(case, 'pendencias-ficha', {'operacao': 'atualizar'})['pendencias']
        for issue in issues:
            if issue['tipo'] != 'comparacao_fontes' or issue['situacao'] == 'resolvida':
                continue
            context = issue['contexto']
            current = context['revisao_atual']
            self.post(case, 'pendencias-ficha', {'operacao': 'resolver', 'client_id': str(uuid4()),
                'pendencia_id': issue['id'], 'assinatura': issue['assinatura'], 'versao_esperada': issue['versao'],
                'acao': 'vincular', 'revisoes_ids': [current['id']],
                'evidencias_ids': [e['id'] for e in context['evidencias']], 'confirmei_revisao': True,
                'justificativa': 'DECISÃO COMPARATIVA SIMULADA: todas as evidências sintéticas revistas, sem homologação profissional'})
        pending = self.client.get(f'/api/v1/atendimentos/{case}/conferencia/').data['pendencias']
        self.assertEqual(pending, [], [(p['chave'], p['tipo'], p['pergunta']) for p in pending])
        return manifest, reference, uncertain_file

    def test_both_full_variants_keep_sources_block_uncertainty_and_export_exact_saved_text(self):
        from docx import Document
        from pypdf import PdfReader
        normalize = lambda text: ' '.join(text.split())
        with patch('openai.OpenAI') as provider:
            for variant in ['direta', 'procuracao']:
                with self.subTest(variant=variant):
                    manifest, reference, file_id = self.full_case(variant)
                    case = manifest['atendimento_id']
                    result = self.generate(case); self.assertEqual(result.status_code, 200, result.data)
                    draft = result.data['versoes'][0]
                    expected_markers = {f'{a}.{k}' for a, k in TOKEN.findall(reference['corpo'])}
                    self.assertEqual(set(draft['dados']), expected_markers)
                    self.assertNotIn('{{', draft['corpo'])
                    self.assertIn('CONSULTA NÃO REALIZADA', draft['corpo'])
                    self.assertIn('SEM RECOLHIMENTO', draft['corpo'])
                    self.assertIn('SEM ASSINATURA', draft['corpo'])
                    self.assertEqual(any(k.startswith('procurador.') for k in draft['dados']), variant == 'procuracao')
                    for marker, datum in draft['dados'].items():
                        alias = marker.split('.')[0]
                        self.assertEqual(datum['objeto_id'], manifest['registros'][alias])
                        self.assertTrue(datum['confirmacao']); self.assertTrue(datum['propostas_ids'])
                        self.assertTrue(datum['origem']['trecho'])
                        if datum['origem']['tipo'] == 'documento':
                            self.assertEqual(datum['origem']['pagina'], 1)
                            self.assertEqual(datum['origem']['leitura_versao'], 1)
                        for occurrence in datum['ocorrencias']:
                            self.assertEqual(draft['corpo'][occurrence['inicio']:occurrence['fim']], datum['valor'])
                    self.assertEqual(self.client.get(f"/api/v1/atendimentos/{case}/minutas/{draft['id']}/exportar/", {'formato': 'pdf'}).status_code, 409)
                    self.post(case, f"minutas/{draft['id']}/revisar", {'versao_esperada': draft['versao'],
                        'conferi_texto': True, 'conferi_dados': True, 'conferi_fontes': True,
                        'motivo': 'REVISÃO FINAL INTEGRAL SIMULADA, SEM HOMOLOGAÇÃO PROFISSIONAL'})
                    for fmt in ['docx', 'pdf']:
                        export = self.client.get(f"/api/v1/atendimentos/{case}/minutas/{draft['id']}/exportar/", {'formato': fmt})
                        self.assertEqual(export.status_code, 200, getattr(export, 'data', None))
                        if fmt == 'docx':
                            self.assertEqual('\n'.join(p.text for p in Document(BytesIO(export.content)).paragraphs[3:]), draft['corpo'])
                        else:
                            text = '\n'.join(p.extract_text() for p in PdfReader(BytesIO(export.content)).pages)
                            for datum in draft['dados'].values():
                                self.assertIn(normalize(datum['valor']), normalize(text))
                    # Reading replacement is deliberately simulated in the test
                    # DB, not reported as OCR or a real external consultation.
                    AnexoAtendimento.objects.filter(id=file_id).update(leitura_versao=2)
                    self.assertEqual(self.generate(case).status_code, 400)
                    self.assertEqual(self.client.get(f"/api/v1/atendimentos/{case}/minutas/{draft['id']}/exportar/", {'formato': 'docx'}).status_code, 409)
            provider.assert_not_called()
