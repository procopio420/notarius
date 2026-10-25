"""Editable Word and paginated PDF from the same saved plain text, with no rewriting."""
from io import BytesIO
from pathlib import Path
from xml.sax.saxutils import escape
from zoneinfo import ZoneInfo
from rest_framework.exceptions import ValidationError

def descricao(minuta):
    date = minuta.created_at.astimezone(ZoneInfo('America/Sao_Paulo')).strftime('%d/%m/%Y %H:%M')
    origin = 'Editada pela escrevente' if minuta.tipo == 'editada' else 'Preenchida pelo sistema'
    return f'Versão {minuta.versao} | Modelo {minuta.modelo.versao} | {origin} | {date}'

def word(minuta):
    from docx import Document
    from docx.shared import Cm, Pt, RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(2.5)
    section.left_margin = section.right_margin = Cm(2.5)
    normal = doc.styles['Normal']
    normal.font.name = 'Times New Roman'
    normal.font.size = Pt(12)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15
    title = doc.add_paragraph('Minuta de compra e venda', style='Title')
    for run in title.runs:
        run.font.color.rgb = RGBColor(0, 0, 0)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(16)
    doc.add_paragraph('Para revisão. Não constitui documento assinado ou ato concluído.')
    doc.add_paragraph(descricao(minuta))
    # Keep each existing line, including blank lines; do not interpret markup or fields.
    for line in minuta.corpo.replace('\r\n', '\n').replace('\r', '\n').split('\n'):
        doc.add_paragraph(line)
    footer = section.footer.paragraphs[0]
    footer.add_run('Minuta para revisão | Página ')
    field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)
    doc.core_properties.author = 'Notarius'
    doc.core_properties.title = 'Minuta de compra e venda para revisão'
    stream = BytesIO(); doc.save(stream)
    return stream.getvalue()

def pdf(minuta):
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import cm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from django.conf import settings
    paths = [getattr(settings, 'NOTARIUS_PDF_FONT', ''), 'C:/Windows/Fonts/times.ttf',
             '/usr/share/fonts/truetype/dejavu/DejaVuSerif.ttf']
    path = next((p for p in paths if p and Path(p).is_file()), None)
    if not path:
        raise ValidationError('A fonte do PDF não está disponível. Configure NOTARIUS_PDF_FONT; o texto não foi substituído.')
    font_name = 'NotariusDocument'
    if font_name not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont(font_name, path))
    font = pdfmetrics.getFont(font_name)
    text = minuta.corpo + descricao(minuta)
    if any(not c.isspace() and ord(c) not in font.face.charToGlyph for c in text):
        raise ValidationError('O texto contém caracteres não suportados pela fonte do PDF. Confira o texto ou use Word; nenhum caractere foi substituído.')
    stream = BytesIO()
    document = SimpleDocTemplate(stream, pagesize=A4, topMargin=2.5*cm, bottomMargin=2.5*cm,
                                 leftMargin=2.5*cm, rightMargin=2.5*cm, title='Minuta de compra e venda para revisão', author='Notarius')
    body = ParagraphStyle('Body', fontName=font_name, fontSize=12, leading=15, spaceAfter=6, alignment=TA_LEFT, splitLongWords=True)
    heading = ParagraphStyle('Heading', parent=body, fontSize=16, leading=20, spaceAfter=12)
    small = ParagraphStyle('Small', parent=body, fontSize=9, leading=12, spaceAfter=12)
    story = [Paragraph('Minuta de compra e venda', heading),
             Paragraph('Para revisão. Não constitui documento assinado ou ato concluído.', small),
             Paragraph(escape(descricao(minuta)), small)]
    for line in minuta.corpo.replace('\r\n', '\n').replace('\r', '\n').split('\n'):
        story.append(Paragraph(escape(line).replace('\t', '&#160;&#160;&#160;&#160;'), body) if line else Spacer(1, 6))
    def footer(canvas, doc):
        canvas.saveState(); canvas.setFont(font_name, 9)
        canvas.drawString(2.5*cm, 1.5*cm, 'Minuta para revisão')
        canvas.drawRightString(A4[0]-2.5*cm, 1.5*cm, f'Página {doc.page}')
        canvas.restoreState()
    document.build(story, onFirstPage=footer, onLaterPages=footer)
    return stream.getvalue()
